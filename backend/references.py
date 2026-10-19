"""Bounded text extraction. Original files are never retained or executed."""
import base64
import binascii
import io
import zipfile
from pathlib import PurePath
from .service import Invalid, text

MAX_FILE=20*1024*1024


def extract(body):
    name=text(body.get('name'),'Имя файла',200,True)
    if PurePath(name).name!=name or '\\' in name:raise Invalid('Укажите имя файла без пути.')
    suffix=PurePath(name).suffix.lower()
    if suffix not in ['.txt','.pdf','.docx']:raise Invalid('Поддерживаются TXT, PDF и DOCX.')
    try:
        encoded=body.get('data','')
        if not isinstance(encoded,str) or len(encoded)>28_000_000:raise ValueError()
        raw=base64.b64decode(encoded,validate=True)
    except (ValueError,binascii.Error):raise Invalid('Некорректное содержимое файла.') from None
    if not raw or len(raw)>MAX_FILE:raise Invalid('Файл пустой или превышает 20 МБ.')
    try:
        if suffix=='.txt':
            content=raw.decode('utf-8-sig')
        elif suffix=='.docx':
            from docx import Document
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if len(archive.infolist())>1000 or sum(x.file_size for x in archive.infolist())>30_000_000:raise Invalid('DOCX слишком большой после распаковки.')
                if archive.getinfo('word/document.xml').file_size>5_000_000:raise Invalid('Текст DOCX слишком большой.')
            doc=Document(io.BytesIO(raw))
            content='\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
        else:
            from pypdf import PdfReader
            if not raw.startswith(b'%PDF-'):raise Invalid('Файл не является PDF.')
            reader=PdfReader(io.BytesIO(raw),strict=True)
            if reader.is_encrypted:raise Invalid('PDF защищён паролем. Загрузите незашифрованный документ.')
            if len(reader.pages)>200:raise Invalid('PDF: допускается до 200 страниц.')
            parts=[]
            for page in reader.pages:
                stream=page.get_contents()
                if stream is not None and len(stream.get_data())>5_000_000:raise Invalid('Содержимое страницы PDF слишком большое.')
                parts.append(page.extract_text() or '')
                if sum(map(len,parts))>60000:raise Invalid('Текст референса превышает 60 000 символов. Выберите нужные разделы.')
            content='\n'.join(parts)
        if not content.strip():raise Invalid('Не найден текст. Для скана PDF нужен OCR; загрузите документ с текстовым слоем.')
        return dict(name=name,text=text(content,'Текст референса',60000,True))
    except Invalid:raise
    except Exception:raise Invalid('Не удалось прочитать файл. Проверьте формат; TXT должен быть в UTF-8.') from None
