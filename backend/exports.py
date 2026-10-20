"""Editable document, lecture slides and Moodle assessment exports."""
import io
import textwrap
from xml.etree import ElementTree as ET
from .curriculum import HOUR_FIELDS


CHECK_LABELS={'materialCounts':'Количество оценочных материалов и слайдов','contentDepth':'Объём учебного текста (не предметная экспертиза)','hours':'Сумма часов','structure':'Полнота каркаса','materials':'Полнота материалов','topicHours':'Часы тем и видов работы','competencyMap':'Карта компетенций','coverage':'Покрытие темами, тестами и заданиями','lecturePlans':'Планы лекций','sources':'Подтверждение источников методистом'}

def source_line(item):
    status='ПОДТВЕРЖДЁН МЕТОДИСТОМ' if item['status']=='confirmed' else 'НЕ ПОДТВЕРЖДЁН — требуется проверка'
    return f"{status}\n{item['title']}\n{item['url']}\n{item['evidence']}"


def outcome_links(course,item):
    return ', '.join(f'РО-{i+1}: {course["competencies"][i]}' for i in item.get('competencyIndices',[])) or 'Связи не заданы'


def docx(course):
    from docx import Document
    from docx.shared import Cm,Pt
    document=Document()
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for border in document.styles['Title'].element.xpath('.//w:pBdr'):border.getparent().remove(border)
    section=document.sections[0];section.page_width=Cm(21);section.page_height=Cm(29.7);section.top_margin=Cm(2);section.bottom_margin=Cm(2);section.left_margin=Cm(2);section.right_margin=Cm(1.5)
    style=document.styles['Normal'];style.font.name='Arial';style.font.size=Pt(10)
    general=course['type']=='Общеразвивающая'
    document.add_heading('Дополнительная общеразвивающая программа' if general else 'Дополнительная профессиональная программа повышения квалификации',0)
    document.add_heading(course['title'],1)
    if course['type']=='Стажировка':document.add_paragraph('Стажировка оформляется как программа повышения квалификации.')
    document.add_paragraph('ЧЕРНОВИК. Утверждает руководитель после проверки методистом и предметным экспертом.')
    document.add_paragraph('Организация: ____________________\nУТВЕРЖДАЮ: ____________________\nПриказ № ____ от __________\nПреподаватель: ____________________\nПодпись: ____________________')
    if course.get('materialsStale'):document.add_paragraph('ВНИМАНИЕ: материалы устарели после изменения каркаса.')
    def heading(title,value=None):
        document.add_heading(title,1)
        if value is not None:document.add_paragraph(str(value) or 'Не заполнено')
    heading('1. Пояснительная записка')
    document.add_paragraph(f"Тип: {course['type']}. Объём: {course['hours']} ч. Форма: {course['format']}. Срок: {course['duration']}.")
    document.add_paragraph(f"Аудитория: {course['audience']}. Входной уровень: {course['level']}.")
    heading('Цель программы',course['goal'])
    if general:document.add_paragraph('Планируемый документ об обучении: ____________________.')
    else:document.add_paragraph('Требования к образованию поступающих, документ о квалификации, действующий профстандарт и коды трудовых функций: заполняет методист после проверки первоисточников.')
    heading('Нормативные основания')
    document.add_paragraph('Действующую редакцию нормативных оснований проверяет методист. Неподтверждённые номера приказов и профстандартов не включаются как установленные факты.')
    heading('2. Результаты обучения')
    for i,value in enumerate(course['competencies']):document.add_paragraph(f'РО-{i+1}. {value}')
    heading('Карта компетенций')
    for item in course.get('competencyMap',[]):
        document.add_heading(f"РО-{item['competencyIndex']+1}",2)
        for key,label in [('knowledge','Знания'),('skills','Умения'),('abilities','Навыки')]:document.add_paragraph(label+': '+'; '.join(item[key]))
        for mi,module in enumerate(course['modules']):
            for ti,topic in enumerate(module.get('topics',[])):
                if item['competencyIndex'] in topic.get('competencyIndices',[]):document.add_paragraph(f"Тема {mi+1}.{ti+1}: {topic['title']}")
        for name,label in [('tests','Вопросы'),('tasks','Задания')]:document.add_paragraph(label+': '+', '.join(str(i+1) for i,t in enumerate(course[name]) if item['competencyIndex'] in t.get('competencyIndices',[])))
    heading('3. Учебный план')
    table=document.add_table(rows=1,cols=2);table.style='Table Grid'
    for cell,value in zip(table.rows[0].cells,['Модуль','Всего часов']):cell.text=value
    for module in course['modules']:
        cells=table.add_row().cells;cells[0].text=module['title'];cells[1].text=str(module['hours'])
    cells=table.add_row().cells;cells[0].text='Итого';cells[1].text=str(sum(m['hours'] for m in course['modules']))
    heading('Учебно-тематический план')
    table=document.add_table(rows=1,cols=7);table.style='Table Grid';table.autofit=False
    widths=[6.7,1.3,1.6,1.8,2,2.3,1.8]
    for column,width in zip(table.columns,widths):column.width=Cm(width)
    header=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(header)
    labels=['Тема','Всего','Лекции','Практика','Лабораторные','Самостоятельная','Аттестация']
    for cell,value in zip(table.rows[0].cells,labels):cell.text=value
    totals=[0]*5
    for mi,module in enumerate(course['modules']):
        for ti,topic in enumerate(module.get('topics',[])):
            values=[topic[k] for k in HOUR_FIELDS];totals=[a+b for a,b in zip(totals,values)]
            row=table.add_row().cells
            for cell,value in zip(row,[f"{mi+1}.{ti+1} {topic['title']}",sum(values),*values]):cell.text=str(value)
    for cell,value in zip(table.add_row().cells,['Итого',sum(totals),*totals]):cell.text=str(value)
    for row in table.rows:
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for cell,width in zip(row.cells,widths):
            cell.width=Cm(width)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:run.font.size=Pt(8)
    heading('Календарный учебный график','Даты и расписание занятий: заполняет методист с учётом срока обучения.')
    heading('4. Рабочие программы модулей')
    for module in course['modules']:
        document.add_heading(module['title'],2);document.add_paragraph(module['content'] or 'Содержание не подготовлено')
    heading('Планы лекций')
    for module in course['modules']:
        document.add_heading(module['title'],2);document.add_paragraph(module.get('lecturePlan') or 'Не подготовлен')
    heading('5. Организационно-педагогические условия','Оборудование, программное обеспечение, кадровое обеспечение и доступность уточняются организатором. Преподаватель: ____________________.')
    heading('6. Формы аттестации и оценочные материалы',course['assessment'])
    document.add_heading('Тестовые вопросы и ключи',2)
    for i,item in enumerate(course['tests']):
        document.add_paragraph(f"{i+1}. {item['question']}\nКлюч: {item['answer']}\n{outcome_links(course,item)}")
    document.add_heading('Задания для семинаров и критерии',2)
    for i,item in enumerate(course['tasks']):document.add_paragraph(f"{i+1}. {item['text']}\nКритерии: {item['criteria']}\n{outcome_links(course,item)}")
    heading('7. Подтверждённые источники')
    confirmed=[s for s in course.get('sources',[]) if s['status']=='confirmed']
    if not confirmed:document.add_paragraph('Подтверждённые источники не добавлены. Литературу и нормативную базу должен проверить методист.')
    for item in confirmed:document.add_paragraph(f"{item['title']}\n{item['url']}\nПодтверждение методиста: {item['evidence']}")
    heading('Неподтверждённые источники — не использовать как проверенную литературу')
    for item in course.get('sources',[]):
        if item['status']!='confirmed':document.add_paragraph(source_line(item))
    heading('Отчёт автоматической проверки')
    from .service import checks
    for key,result in checks(course).items():
        document.add_paragraph(f"{CHECK_LABELS.get(key,key)}: {'пройдено' if result['valid'] else 'требует проверки/исправления'}")
        for detail in result.get('details',[]):document.add_paragraph(detail)
    document.add_paragraph('Проверка структуры не подтверждает предметную корректность. Доступность внешних ссылок проверяется отдельно по запросу методиста и не подтверждает достоверность источников.')
    buffer=io.BytesIO();document.save(buffer);return buffer.getvalue()


def pptx(course):
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.dml.color import RGBColor
    from .service import Invalid
    if not course['slides']:raise Invalid('Слайды ещё не подготовлены.')
    exported_slides=list(course['slides'])
    source_lines=[source_line(item) for item in course.get('sources',[])] or ['Источники не добавлены. Литература требует проверки методистом.']
    exported_slides.extend(dict(title='Источники и статус проверки',body=line,moduleIndex=0,appendix=True) for line in source_lines)
    from .service import checks
    exported_slides.extend(dict(title='Отчёт автоматической проверки',body=f"{CHECK_LABELS.get(key,key)}: {'пройдено' if result['valid'] else 'требует проверки/исправления'}"+('\n'+'\n'.join(result.get('details',[])) if result.get('details') else ''),moduleIndex=0,appendix=True) for key,result in checks(course).items())
    deck=Presentation();deck.slide_width=Inches(13.333);deck.slide_height=Inches(7.5)
    for item in exported_slides:
        chunks=textwrap.wrap(item['body'],width=550,break_long_words=True,replace_whitespace=False) or ['']
        for part,body in enumerate(chunks):
            slide=deck.slides.add_slide(deck.slide_layouts[6])
            def box(x,y,w,h,value,size,color=(24,39,32)):
                frame=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
                frame.word_wrap=True
                for i,line in enumerate(value.splitlines() or ['']):
                    p=frame.paragraphs[0] if i==0 else frame.add_paragraph();p.text=line;p.font.name='Arial';p.font.size=Pt(size);p.font.color.rgb=RGBColor(*color)
                return frame
            box(.7,.25,11.9,.5,course['title'],13,(0,124,61))
            box(.7,1,11.9,1.3,item['title']+(f' · {part+1}/{len(chunks)}' if len(chunks)>1 else ''),30)
            box(.7,2.5,11.9,3.9,body,22)
            box(.7,6.8,11.9,.4,f"Черновик · {'Проверки и источники' if item.get('appendix') else 'Модуль '+str(item['moduleIndex']+1)} · Слайд {len(deck.slides)}",12)
            slide.notes_slide.notes_text_frame.text='Проверка структуры и объёма не подтверждает предметную корректность. Доступность ссылок проверяется отдельно по запросу методиста; её временный результат не включён в этот экспорт.' if item.get('appendix') else course['modules'][item['moduleIndex']].get('lecturePlan','')
    buffer=io.BytesIO();deck.save(buffer);return buffer.getvalue()


def moodle(course):
    from .service import Invalid
    if not course['tests']:raise Invalid('Тесты ещё не подготовлены.')
    root=ET.Element('quiz')
    for i,item in enumerate(course['tests']):
        question=ET.SubElement(root,'question',type='essay')
        ET.SubElement(ET.SubElement(question,'name'),'text').text=f'{course["title"]}: вопрос {i+1}'
        ET.SubElement(ET.SubElement(question,'questiontext',format='plain_text'),'text').text=item['question']
        ET.SubElement(ET.SubElement(question,'generalfeedback',format='plain_text'),'text').text=outcome_links(course,item)
        ET.SubElement(question,'defaultgrade').text='1'
        ET.SubElement(question,'responseformat').text='editor'
        ET.SubElement(question,'responserequired').text='1'
        # Moodle essay keys belong to graderinfo, never learner-visible feedback.
        ET.SubElement(ET.SubElement(question,'graderinfo',format='plain_text'),'text').text=item['answer']
    return ET.tostring(root,encoding='utf-8',xml_declaration=True)
