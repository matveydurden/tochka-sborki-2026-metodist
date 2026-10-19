import io
import unittest
from backend import service
from test_curriculum import curriculum
from test_service import materials

class ExportTests(unittest.TestCase):
    def course(self):return service.normalize(dict(curriculum(),**materials()))
    def test_docx_has_different_type_sections_hours_and_outcome_links(self):
        from docx import Document
        data=service.export_docx(self.course())
        doc=Document(io.BytesIO(data));text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        for value in ['Учебно-тематический план','Лекции','Самостоятельная','Карта компетенций','Планы лекций','Подтверждённые источники']:
            self.assertIn(value,text)
        course=self.course();course['type']='Общеразвивающая'
        doc=Document(io.BytesIO(service.export_docx(course)))
        self.assertIn('Дополнительная общеразвивающая программа','\n'.join(p.text for p in doc.paragraphs))
    def test_pptx_reopens_and_keeps_all_text_without_truncation(self):
        from backend.exports import pptx
        from pptx import Presentation
        course=self.course();course['slides'][0]['body']='Текст & пример <код> '*120
        deck=Presentation(io.BytesIO(pptx(course)))
        self.assertGreater(len(deck.slides),len(course['slides']))
        text=' '.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
        self.assertEqual(text.count('Текст & пример <код>'),120)
    def test_moodle_xml_preserves_escaped_questions_and_keys(self):
        from backend.exports import moodle
        from xml.etree import ElementTree as ET
        course=self.course();course['tests'][0]['question']='Что <tag> & "x"?'
        root=ET.fromstring(moodle(course))
        self.assertEqual(root[0].find('questiontext/text').text,course['tests'][0]['question'])
        self.assertEqual(root[0].find('graderinfo/text').text,course['tests'][0]['answer'])
