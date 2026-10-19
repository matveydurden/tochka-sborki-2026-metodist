import base64
import io
import unittest
import zipfile
from backend import service

class FileTests(unittest.TestCase):
    def extract(self,name,data):
        from backend.references import extract
        return extract(dict(name=name,data=base64.b64encode(data).decode()))

    def test_utf8_and_docx_text_is_extracted_without_saving_raw_file(self):
        self.assertEqual(self.extract('Требования.txt','Нужно 24 часа'.encode())['text'],'Нужно 24 часа')
        from docx import Document
        doc=Document();doc.add_paragraph('Практический курс');doc.add_table(rows=1,cols=1).cell(0,0).text='24 часа'
        buffer=io.BytesIO();doc.save(buffer)
        text=self.extract('Пример.docx',buffer.getvalue())['text']
        self.assertIn('Практический курс',text);self.assertIn('24 часа',text)

    def test_corrupt_unsupported_and_binary_inputs_do_not_become_references(self):
        for name,data in [('bad.docx',b'not a zip'),('bad.pdf',b'not a pdf'),('bad.txt',b'\x00\x01'),('run.exe',b'text')]:
            with self.subTest(name=name),self.assertRaises(service.Invalid):self.extract(name,data)

    def test_pdf_text_and_scanned_pdf_are_distinguished(self):
        from pypdf import PdfWriter
        pdf=PdfWriter();pdf.add_blank_page(width=200,height=200)
        buffer=io.BytesIO();pdf.write(buffer)
        with self.assertRaisesRegex(service.Invalid,'текст|OCR'):self.extract('scan.pdf',buffer.getvalue())
